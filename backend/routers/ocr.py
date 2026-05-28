import logging
from io import BytesIO
from typing import List, Optional, Literal
from fastapi import APIRouter, HTTPException
from pydantic import BaseModel
import fitz
from PIL import Image
from utils.session_manager import SessionManager
from utils import ai_clients

router = APIRouter()
session_manager = SessionManager()
logger = logging.getLogger(__name__)

class OCRRequest(BaseModel):
    session_id: str
    filename: str
    output_format: Literal["docx", "xlsx"]
    lang: str = "eng"
    ai_provider: Optional[str] = None
    ai_api_key: Optional[str] = None
    ai_model: Optional[str] = None
    pages: Optional[List[int]] = None
    output_name: Optional[str] = None

@router.post("/ocr")
async def ocr_pdf(req: OCRRequest):
    try:
        session_dir = session_manager.get_session_path(req.session_id)
        file_path = session_manager.get_file_path(req.session_id, req.filename)
        base_name = req.filename.rsplit(".", 1)[0]
        output_name = req.output_name or f"{base_name}_ocr.{req.output_format}"
        if not output_name.endswith(f".{req.output_format}"):
            output_name += f".{req.output_format}"
        output_path = session_dir / output_name
        use_ai = bool(req.ai_provider and req.ai_api_key)
        mat = fitz.Matrix(300 / 72, 300 / 72)
        page_texts = []
        with fitz.open(str(file_path)) as doc:
            total = doc.page_count
            target_pages = req.pages if req.pages else list(range(1, total + 1))
            for p in target_pages:
                if p < 1 or p > total:
                    raise HTTPException(status_code=400, detail=f"Page {p} out of range")
                pix = doc[p - 1].get_pixmap(matrix=mat)
                png_bytes = pix.tobytes("png")
                if use_ai:
                    text = await ai_clients.get_ocr_result(png_bytes, req.ai_provider, req.ai_api_key, req.ai_model)
                else:
                    import pytesseract
                    img = Image.open(BytesIO(png_bytes))
                    text = pytesseract.image_to_string(img, lang=req.lang)
                page_texts.append((p, text))
        if req.output_format == "docx":
            from docx import Document
            word_doc = Document()
            word_doc.add_heading(base_name, 0)
            for page_num, text in page_texts:
                word_doc.add_heading(f"Page {page_num}", level=1)
                for para in text.split("\n\n"):
                    if para.strip():
                        word_doc.add_paragraph(para.strip())
            word_doc.save(str(output_path))
        else:
            from openpyxl import Workbook
            wb = Workbook()
            wb.remove(wb.active)
            for page_num, text in page_texts:
                ws = wb.create_sheet(title=f"Page {page_num}")
                for row_idx, line in enumerate(text.split("\n"), start=1):
                    ws.cell(row=row_idx, column=1, value=line)
            wb.save(str(output_path))
        return {
            "output_filename": output_name,
            "pages_processed": len(page_texts),
            "method": "ai" if use_ai else "local",
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error("OCR error: %s", e)
        raise HTTPException(status_code=500, detail=str(e))
